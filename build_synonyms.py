"""
Build a synonym database by lemmatizing words from ordliste Excel file.
Uses spaCy Norwegian model to group word variants by their lemma (stem).
"""

from pathlib import Path
import sqlite3
import sys
from collections import defaultdict

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

try:
    import spacy
except ImportError:
    print("ERROR: spaCy not installed. Run: pip install spacy")
    sys.exit(1)


EXCEL_PATH = Path.cwd() / "Ordliste_Norsk_ny.xlsx"
ORDLISTE_SHEET_CANDIDATES = ["Ordliste", "AlleOrd", "Kolonner"]
COLUMN_START = "A"
COLUMN_END = "AC"
SYNONYMS_DB = Path(__file__).with_name("ordliste_synonyms.sqlite3")


def normalize_word(value: object) -> str:
    """Normalize word for comparison."""
    if value is None:
        return ""
    text = str(value).strip().lower()
    return " ".join(text.split())


def load_spacy_model():
    """Load Norwegian spaCy model, download if needed."""
    try:
        nlp = spacy.load("nb_core_news_sm")
        print("✓ Loaded spaCy Norwegian model (nb_core_news_sm)")
        return nlp
    except OSError:
        print("Downloading Norwegian spaCy model...")
        import subprocess
        subprocess.check_call([sys.executable, "-m", "spacy", "download", "nb_core_news_sm"])
        nlp = spacy.load("nb_core_news_sm")
        print("✓ Downloaded and loaded Norwegian model")
        return nlp


def read_words_from_excel(excel_path: Path) -> list[str]:
    """Read all words from Excel ordliste."""
    if not excel_path.is_file():
        raise FileNotFoundError(f"Excel file not found: {excel_path}")

    workbook = load_workbook(
        excel_path,
        read_only=True,
        data_only=True,
        keep_vba=True,
        keep_links=False,
    )
    try:
        # Find the correct sheet
        sheet = None
        for candidate in ORDLISTE_SHEET_CANDIDATES:
            for worksheet in workbook.sheetnames:
                if worksheet.strip().casefold() == candidate.strip().casefold():
                    sheet = workbook[worksheet]
                    break
            if sheet is not None:
                break

        if sheet is None:
            available = ", ".join(workbook.sheetnames)
            raise KeyError(f"Sheet not found. Available: {available}")

        # Extract words from columns A-AC
        words = []
        min_column = column_index_from_string(COLUMN_START)
        max_column = column_index_from_string(COLUMN_END)

        for row_values in sheet.iter_rows(min_col=min_column, max_col=max_column, values_only=True):
            for cell_value in row_values:
                word = normalize_word(cell_value)
                if word:
                    words.append(word)

        return words
    finally:
        workbook.close()


def build_synonym_groups(words: list[str], nlp) -> dict[str, list[str]]:
    """Group words by lemma using spaCy."""
    groups = defaultdict(list)
    processed = 0
    skipped = 0

    for word in words:
        processed += 1
        if processed % 100000 == 0:
            print(f"  Processed {processed:,} words...".replace(",", " "))

        try:
            # Lemmatize the word
            doc = nlp(word)
            if doc:
                lemma = doc[0].lemma_.lower().strip()
                if lemma:
                    groups[lemma].append(word)
                else:
                    skipped += 1
            else:
                skipped += 1
        except Exception:
            skipped += 1

    print(f"Processed {processed:,} words, {skipped:,} skipped".replace(",", " "))
    return groups


def setup_database(db_path: Path) -> sqlite3.Connection:
    """Create or open synonyms database."""
    connection = sqlite3.connect(db_path)
    connection.execute("PRAGMA journal_mode=WAL")
    connection.execute("PRAGMA synchronous=NORMAL")

    # Create tables
    connection.execute("""
        CREATE TABLE IF NOT EXISTS lemmas (
            lemma TEXT PRIMARY KEY,
            variant_count INTEGER NOT NULL
        )
    """)
    connection.execute("""
        CREATE TABLE IF NOT EXISTS variants (
            variant TEXT PRIMARY KEY,
            lemma TEXT NOT NULL,
            FOREIGN KEY (lemma) REFERENCES lemmas(lemma)
        )
    """)
    connection.execute("""
        CREATE INDEX IF NOT EXISTS idx_variants_lemma ON variants(lemma)
    """)
    connection.commit()
    return connection


def save_synonym_groups(connection: sqlite3.Connection, groups: dict[str, list[str]]) -> None:
    """Save lemma groups to database."""
    with connection:
        # Clear existing data
        connection.execute("DELETE FROM variants")
        connection.execute("DELETE FROM lemmas")

        # Insert lemmas and variants
        for lemma, variants in groups.items():
            unique_variants = list(set(variants))  # Remove duplicates
            if len(unique_variants) > 0:
                connection.execute(
                    "INSERT INTO lemmas (lemma, variant_count) VALUES (?, ?)",
                    (lemma, len(unique_variants)),
                )
                for variant in unique_variants:
                    connection.execute(
                        "INSERT INTO variants (variant, lemma) VALUES (?, ?)",
                        (variant, lemma),
                    )

        variant_count = connection.execute("SELECT COUNT(*) FROM variants").fetchone()[0]
        lemma_count = connection.execute("SELECT COUNT(*) FROM lemmas").fetchone()[0]

    print(f"✓ Saved {lemma_count:,} lemmas and {variant_count:,} variants".replace(",", " "))


def main() -> int:
    print("Building Norwegian synonym database...")
    print(f"Excel file: {EXCEL_PATH}")
    print(f"Database: {SYNONYMS_DB}")
    print()

    # Load spaCy model
    print("Loading spaCy Norwegian model...")
    nlp = load_spacy_model()
    print()

    # Read words
    print("Reading words from Excel...")
    try:
        words = read_words_from_excel(EXCEL_PATH)
        print(f"✓ Read {len(words):,} words".replace(",", " "))
    except Exception as e:
        print(f"ERROR: {e}")
        return 1
    print()

    # Build groups
    print("Lemmatizing and grouping words...")
    groups = build_synonym_groups(words, nlp)
    print(f"✓ Found {len(groups):,} unique lemmas".replace(",", " "))
    print()

    # Save to database
    print("Saving to database...")
    try:
        connection = setup_database(SYNONYMS_DB)
        save_synonym_groups(connection, groups)
        connection.close()
    except Exception as e:
        print(f"ERROR: {e}")
        return 1

    print()
    print("Done! You can now query the synonym database.")
    print("Example queries:")
    print("  SELECT lemma, GROUP_CONCAT(variant, ', ') FROM variants GROUP BY lemma LIMIT 5;")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
