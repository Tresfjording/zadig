import argparse
import re
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
    from openpyxl.utils import range_boundaries, get_column_letter
except ImportError:
    load_workbook = None
    range_boundaries = None
    get_column_letter = None

CELL_REF_RE = re.compile(
    r"(?:(?:'(?P<sheet_quoted>[^']+)')|(?P<sheet_unquoted>[A-Za-z0-9_]+))?!"
    r"\$?(?P<col>[A-Z]{1,3})\$?(?P<row>\d+)"
    r"(?:\s*:\s*\$?(?P<col2>[A-Z]{1,3})\$?(?P<row2>\d+))?",
    re.IGNORECASE,
)

CELL_REF_SHORT_RE = re.compile(r"\$?(?P<col>[A-Z]{1,3})\$?(?P<row>\d+)", re.IGNORECASE)


def normalize_sheet_name(sheet_name: str) -> str:
    return sheet_name.strip().replace("'", "")


def expand_cell_reference(sheet: str, start_col: str, start_row: str, end_col: str = None, end_row: str = None) -> set[str]:
    if end_col and end_row:
        if range_boundaries is None or get_column_letter is None:
            return {f"{sheet}!{start_col.upper()}{start_row}"}
        min_col, min_row, max_col, max_row = range_boundaries(f"{start_col}{start_row}:{end_col}{end_row}")
        cells = set()
        for row in range(min_row, max_row + 1):
            for col in range(min_col, max_col + 1):
                cells.add(f"{sheet}!{get_column_letter(col)}{row}")
        return cells
    return {f"{sheet}!{start_col.upper()}{start_row}"}


def parse_references(formula: str, current_sheet: str) -> set[str]:
    refs = set()
    for match in CELL_REF_RE.finditer(formula):
        sheet = match.group("sheet_quoted") or match.group("sheet_unquoted") or current_sheet
        sheet = normalize_sheet_name(sheet)
        refs.update(expand_cell_reference(
            sheet,
            match.group("col"),
            match.group("row"),
            match.group("col2"),
            match.group("row2"),
        ))

    if not refs:
        sheet = normalize_sheet_name(current_sheet)
        for match in CELL_REF_SHORT_RE.finditer(formula):
            refs.add(f"{sheet}!{match.group('col').upper()}{match.group('row')}")
    return refs


def build_dependency_graph(workbook) -> dict[str, set[str]]:
    graph: dict[str, set[str]] = {}
    sheet_names = {ws.title for ws in workbook.worksheets}

    for ws in workbook.worksheets:
        for row in ws.iter_rows(values_only=False):
            for cell in row:
                if cell.data_type == "f" and cell.value:
                    node = f"{ws.title}!{cell.coordinate}"
                    formula = str(cell.value)
                    refs = parse_references(formula, ws.title)
                    refs = {ref for ref in refs if ref.split("!")[0] in sheet_names}
                    graph[node] = refs
    return graph


def find_cycles(graph: dict[str, set[str]]) -> list[list[str]]:
    visited: set[str] = set()
    stack: list[str] = []
    cycles: list[list[str]] = []

    def dfs(node: str) -> None:
        if node in stack:
            cycle_start = stack.index(node)
            cycles.append(stack[cycle_start:] + [node])
            return
        if node in visited:
            return
        visited.add(node)
        stack.append(node)
        for neighbour in graph.get(node, set()):
            if neighbour in graph:
                dfs(neighbour)
        stack.pop()

    for node in graph:
        if node not in visited:
            dfs(node)
    return cycles


def check_circular_reference_openpyxl(file_path: Path) -> bool:
    if load_workbook is None:
        raise RuntimeError("openpyxl er ikke installert. Installer med pip install openpyxl.")

    workbook = load_workbook(file_path, data_only=False)
    graph = build_dependency_graph(workbook)
    if not graph:
        print("Ingen formler funnet i arbeidsboken.")
        return False

    cycles = find_cycles(graph)
    if not cycles:
        print("Ingen sirkulære formler funnet via openpyxl.")
        return False

    print("Sirkulære formler funnet via openpyxl:")
    for idx, cycle in enumerate(cycles, start=1):
        print(f"Cycle {idx}:")
        print("  " + " -> ".join(cycle))
    return True


def main() -> int:
    parser = argparse.ArgumentParser(
        description="Sjekk for sirkulære formler i en Excel-arbeidsbok."
    )
    parser.add_argument("fil", type=Path, help="Sti til Excel-filen (.xlsx eller .xlsm)")
    parser.add_argument(
        "--mode",
        choices=["openpyxl"],
        default="openpyxl",
        help="Skannemetode: 'openpyxl' leser filen direkte (anbefalt og eneste støttede).",
    )
    args = parser.parse_args()

    if not args.fil.exists():
        print(f"FEIL: Filen finnes ikke: {args.fil}")
        return 1

    if load_workbook is None:
        print("FEIL: openpyxl er ikke installert.")
        return 1
    found = check_circular_reference_openpyxl(args.fil)

    return 0 if not found else 2


if __name__ == "__main__":
    raise SystemExit(main())
